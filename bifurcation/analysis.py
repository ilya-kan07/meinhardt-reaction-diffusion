# analysis.py
import numpy as np
import matplotlib.pyplot as plt
import math


def find_intersections_general(params, n_points=25000, tol=5e-4):
    """Поиск пересечений a1 и a2 для текущих параметров"""
    c = params['c']
    mu = params['mu']
    eps = params['eps']
    nu = params['nu']
    e = params['e']
    eta = params['eta']
    d = params['d']
    f = params['f']
    c0 = params['c0']

    def a2(y):
        return y * (e + e * f * y**2 - eta * y) / (d * (1 + f * y**2))

    def discriminant(y):
        return c0**2 * c**2 - 4 * c * mu**2 * (eps * y + nu)

    def a1_plus(y):
        disc = discriminant(y)
        return np.where(disc >= 0,
                        (c0 * c + np.sqrt(np.maximum(disc, 0))) / (2 * c * mu),
                        np.nan)

    def a1_minus(y):
        disc = discriminant(y)
        return np.where(disc >= 0,
                        (c0 * c - np.sqrt(np.maximum(disc, 0))) / (2 * c * mu),
                        np.nan)

    y = np.linspace(-3.0, 3.0, n_points)
    ap = a1_plus(y)
    am = a1_minus(y)
    a_two = a2(y)

    crossings = []
    for a1 in [ap, am]:
        valid = ~np.isnan(a1)
        i = 0
        while i < len(y) - 1:
            if valid[i] and valid[i + 1]:
                d1 = a1[i] - a_two[i]
                d2 = a1[i + 1] - a_two[i + 1]
                if d1 * d2 < 0 or (abs(d1) < tol and abs(d2) < tol):
                    denom = d1 - d2
                    frac = d1 / denom if abs(denom) > 1e-12 else 0.5
                    y_cross = y[i] + frac * (y[i + 1] - y[i])
                    a_cross = a_two[i] + frac * (a_two[i + 1] - a_two[i])
                    crossings.append((y_cross, a_cross))
                    i += 20
                    continue
            i += 1

    # Удаляем дубликаты
    unique = []
    for cr in crossings:
        if not any(abs(cr[0] - u[0]) < 1e-5 for u in unique):
            unique.append(cr)
    unique.sort(key=lambda x: x[0])
    return unique

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def plot_a1_a2_intersections(params, varied_param, start, end, num_points=5, ax=None, canvas=None, log_func=None, save_excel=True, excel_filename=None):
    """
    Исправленная версия с надёжной очисткой фигуры при работе внутри GUI.
    Теперь сохраняет координаты равновесия P в Excel файл.
    """
    values = np.linspace(start, end, num_points)

    if ax is None:
        fig, ax = plt.subplots(figsize=(15, 9))
        external = True
    else:
        external = False
        # === КРИТИЧНОЕ ИСПРАВЛЕНИЕ ===
        fig = ax.figure
        fig.clf()                    # полностью очищаем фигуру
        ax = fig.add_subplot(111)    # создаём новый единственный ax

    colors = ['#1f77b4', '#ff7f0e', '#2ca02c', '#d62728', '#9467bd', '#8c564b']

    # === НОВОЕ: Собираем данные для Excel ===
    excel_data = []
    all_intersections_data = []  # для всех пересечений

    # Заголовок в лог
    if log_func:
        log_func(
            f"=== Анализ пересечений a1(y) и a2(y) при изменении {varied_param}\n")
        log_func(f" Параметр     | Кол-во пересечений | Равновесие P (y, a)")
        log_func("-" * 55)
    else:
        print(
            f"\n=== Анализ пересечений a1(y) и a2(y) при изменении {varied_param} ===\n")
        print(f"{' ':<12} | Кол-во пересечений | Основная P (y, a)")
        print("-" * 60)

    for idx, val in enumerate(values):
        temp_params = params.copy()
        temp_params[varied_param] = val

        intersections = find_intersections_general(temp_params)

        if len(intersections) == 6:
            main_y, main_a = intersections[5]
            main_info = f"({main_y:.6f}, {main_a:.6f})"
            has_main_P = True
            p_index = 5
        elif len(intersections) == 5:
            main_y, main_a = intersections[4]
            main_info = f"({main_y:.6f}, {main_a:.6f}) *"
            has_main_P = True
            p_index = 4
        elif len(intersections) == 4:
            main_y, main_a = intersections[3]
            main_info = f"({main_y:.6f}, {main_a:.6f}) *"
            has_main_P = True
            p_index = 3
        else:
            main_y = main_a = None
            main_info = "—"
            has_main_P = False
            p_index = None

        num_inter = len(intersections)

        log_line = f"{varied_param} = {val:8.5f} |         {num_inter:>2}         | {main_info}"

        if log_func:
            log_func(log_line)
        else:
            print(log_line)

        # === НОВОЕ: Сохраняем данные для Excel ===
        excel_data.append({
            'Параметр': varied_param,
            'Значение_параметра': val,
            'Количество_пересечений': num_inter,
            'P_y': main_y if main_y is not None else np.nan,
            'P_a': main_a if main_a is not None else np.nan,
            'P_индекс': p_index if p_index is not None else np.nan,
            'P_статус': 'Найдено' if has_main_P else 'Не найдено'
        })

        # Сохраняем все пересечения для каждого значения параметра
        for inter_idx, (y_cross, a_cross) in enumerate(intersections):
            all_intersections_data.append({
                'Параметр': varied_param,
                'Значение_параметра': val,
                'Номер_пересечения': inter_idx + 1,
                'y': y_cross,
                'a': a_cross,
                'Это_точка_P': (inter_idx == p_index) if p_index is not None else False
            })

        y_plot = np.linspace(-1.8, 4.0, 4000)

        # a2(y) — рисуем только один раз
        if idx == 0:
            def a2_func(y):
                return y * (temp_params['e'] + temp_params['e']*temp_params['f']*y**2 -
                            temp_params['eta']*y) / (temp_params['d'] * (1 + temp_params['f']*y**2))
            ax.plot(y_plot, a2_func(y_plot), color='blue',
                    linewidth=3.5, label=r'$a_2(y)$')

        # a1+ и a1-
        def a1_plus_func(y):
            disc = temp_params['c0']**2 * temp_params['c']**2 - \
                4 * temp_params['c'] * temp_params['mu']**2 * \
                (temp_params['eps']*y + temp_params['nu'])
            return np.where(disc >= 0,
                            (temp_params['c0']*temp_params['c'] + np.sqrt(np.maximum(disc, 0))) /
                            (2*temp_params['c']*temp_params['mu']), np.nan)

        def a1_minus_func(y):
            disc = temp_params['c0']**2 * temp_params['c']**2 - \
                4 * temp_params['c'] * temp_params['mu']**2 * \
                (temp_params['eps']*y + temp_params['nu'])
            return np.where(disc >= 0,
                            (temp_params['c0']*temp_params['c'] - np.sqrt(np.maximum(disc, 0))) /
                            (2*temp_params['c']*temp_params['mu']), np.nan)

        color = colors[idx % len(colors)]
        label = f'{varied_param} = {val:.4f}  ({num_inter} пересеч.)'

        ap = a1_plus_func(y_plot)
        am = a1_minus_func(y_plot)

        ax.plot(y_plot, ap, color=color, linewidth=2.2, label=label)
        ax.plot(y_plot, am, color=color, linewidth=2.2)

        # Точки пересечений
        for y_cross, a_cross in intersections:
            ax.plot(y_cross, a_cross, 'o', color=color, markersize=6,
                    markeredgecolor='black', markeredgewidth=0.8, alpha=0.85)

        # Главная точка P
        if has_main_P and main_y is not None and main_a is not None:
            ax.plot(main_y, main_a, 'o', color=color, markersize=13,
                    markeredgecolor='black', markeredgewidth=2.5)
            ax.text(main_y + 0.05, main_a + 0.16, 'P', fontsize=14,
                    fontweight='bold', color='black', ha='left')

    ax.set_xlabel('$y$', fontsize=14)
    ax.set_ylabel('$a$', fontsize=14)
    ax.set_title(
        f'Влияние параметра ${varied_param}$ на пересечения $a_1(y)$ и $a_2(y)$',
        fontsize=15, pad=20)

    ax.grid(True, alpha=0.35)
    ax.legend(fontsize=10.5, loc='upper right')
    ax.set_xlim(-1.8, 4.0)
    ax.set_ylim(-1.2, 3.2)

    if external:
        plt.tight_layout()
        plt.show()
    else:
        canvas.draw()

    # === НОВОЕ: Сохраняем в Excel ===
    if save_excel:
        if excel_filename is None:
            from datetime import datetime
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            excel_filename = f"equilibrium_points_{varied_param}_{timestamp}.xlsx"

        # Создаем Excel файл с несколькими листами
        with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
            # Лист 1: Основные точки равновесия P
            df_main = pd.DataFrame(excel_data)
            df_main.to_excel(writer, sheet_name='Точки_P', index=False)

            # Лист 2: Все пересечения
            df_all = pd.DataFrame(all_intersections_data)
            df_all.to_excel(writer, sheet_name='Все_пересечения', index=False)

            # Лист 3: Параметры модели
            params_df = pd.DataFrame([
                {'Параметр': k, 'Значение': v}
                for k, v in params.items()
            ])
            params_df.to_excel(writer, sheet_name='Параметры', index=False)

            # Лист 4: Информация об анализе
            info_df = pd.DataFrame({
                'Метрика': ['Изменяемый_параметр', 'Диапазон_от', 'Диапазон_до', 'Количество_точек', 'Формат_P'],
                'Значение': [varied_param, start, end, num_points, 'y, a']
            })
            info_df.to_excel(writer, sheet_name='Информация', index=False)

        # Стилизация Excel файла
        from openpyxl import load_workbook

        wb = load_workbook(excel_filename)

        # Стилизация каждого листа
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Заголовки жирным шрифтом
            header_font = Font(bold=True, size=11)
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font_white = Font(bold=True, size=11, color="FFFFFF")

            for cell in ws[1]:
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')

            # Автоширина колонок
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 30)
                ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(excel_filename)

        if log_func:
            log_func(f"\n✓ Данные сохранены в файл: {excel_filename}")
        else:
            print(f"\n✓ Данные сохранены в файл: {excel_filename}")

    return excel_data if save_excel else None

# ====================== УНИВЕРСАЛЬНЫЙ МЕТОД: ЗАВИСИМОСТЬ КОЭФФИЦИЕНТОВ b1, b3, B ======================


def find_point_P_general(c0, params):
    """Универсальная функция поиска точки P (y, a, s)"""
    def a2(y):
        return y * (params['e'] + params['e'] * params['f'] * y**2 -
                    params['eta'] * y) / (params['d'] * (1 + params['f'] * y**2))

    def discriminant(y):
        return (c0**2 * params['c']**2 -
                4 * params['c'] * params['mu']**2 * (params['eps'] * y + params['nu']))

    def a1_plus(y):
        disc = discriminant(y)
        return np.where(disc >= 0,
                        (c0 * params['c'] + np.sqrt(np.maximum(disc, 0))) /
                        (2 * params['c'] * params['mu']), np.nan)

    def a1_minus(y):
        disc = discriminant(y)
        return np.where(disc >= 0,
                        (c0 * params['c'] - np.sqrt(np.maximum(disc, 0))) /
                        (2 * params['c'] * params['mu']), np.nan)

    y_arr = np.linspace(-3.0, 3.0, 25000)
    ap = a1_plus(y_arr)
    am = a1_minus(y_arr)
    a_two = a2(y_arr)

    crossings = []
    tol = 5e-4

    for a1 in [ap, am]:
        valid = ~np.isnan(a1)
        i = 0
        while i < len(y_arr) - 1:
            if valid[i] and valid[i + 1]:
                d1 = a1[i] - a_two[i]
                d2 = a1[i + 1] - a_two[i + 1]
                if d1 * d2 < 0 or (abs(d1) < tol and abs(d2) < tol):
                    denom = d1 - d2
                    frac = d1 / denom if abs(denom) > 1e-12 else 0.5
                    y_cross = y_arr[i] + frac * (y_arr[i + 1] - y_arr[i])
                    a_cross = a_two[i] + frac * (a_two[i + 1] - a_two[i])
                    crossings.append((y_cross, a_cross))
                    i += 20
                    continue
            i += 1

    unique = []
    for cr in crossings:
        if not any(abs(cr[0] - u[0]) < 1e-5 for u in unique):
            unique.append(cr)
    unique.sort(key=lambda x: x[0])

    if len(unique) == 6:
        y_p, a_p = unique[5]
    elif len(unique) == 5:
        y_p, a_p = unique[4]
    elif len(unique) == 4:
        y_p, a_p = unique[3]
    else:
        return None

    s_p = params['mu'] / (params['c'] * a_p)
    return y_p, a_p, s_p


def zero_stability_coefficients_with_k(param_value, params, varied_param, k=0, L=1.0):
    """Вычисление b1, b3, B"""
    temp_params = params.copy()

    if varied_param == 'c0':
        c0 = param_value
    else:
        temp_params[varied_param] = param_value
        c0 = params.get('c0', 1.128)

    result = find_point_P_general(c0, temp_params)
    if result is None:
        return None, None, None

    y, a, s = result

    c = temp_params['c']
    mu = temp_params['mu']
    nu = temp_params['nu']
    eps = temp_params['eps']
    d = temp_params['d']
    e = temp_params['e']
    eta = temp_params['eta']
    f = temp_params['f']
    Da = temp_params.get('Da', 0.062)
    Ds = temp_params.get('Ds', 1.0)
    Dy = temp_params.get('Dy', 1.0)

    kappa2 = (math.pi * k / L) ** 2

    # Якобиан при k=0
    J11 = 2*c*a*s - mu
    J12 = c*a**2
    J21 = -2*c*a*s
    J22 = -c*a**2 - nu - eps*y
    J23 = -eps*s
    J31 = d
    J33 = -e + 2*eta*y / (1 + f*y**2)**2

    # Коэффициенты для k=0 (не используются в графике, оставлены для справки)
    b1 = -(J11 + J22 + J33)
    b2 = J11*J22 - J12*J21 + J22*J33 + J11*J33
    b3 = -(J11*J22*J33 - J12*J21*J33 + J12*J23*J31)
    B = b1*b2 - b3

    # С диффузией (для выбранной гармоники k)
    Jk11 = J11 - kappa2 * Da
    Jk22 = J22 - kappa2 * Ds
    Jk33 = J33 - kappa2 * Dy

    bk1 = -(Jk11 + Jk22 + Jk33)
    bk2 = Jk11*Jk22 - J12*J21 + Jk22*Jk33 + Jk11*Jk33
    bk3 = -(Jk11*Jk22*Jk33 - J12*J21*Jk33 + J12*J23*J31)
    Bk = bk1 * bk2 - bk3

    return bk1, bk3, Bk


def plot_b_coefficients_vs_parameter(params, varied_param, start, end, num_points=100,
                                     k=0, L=1.0,
                                     ax=None, canvas=None, log_func=None):
    """Построение графиков bk1, bk3, Bk с красивой разметкой начала и конца диапазона"""
    values = np.linspace(start, end, num_points)
    bk1_list, bk3_list, Bk_list, valid_values = [], [], [], []

    if log_func:
        log_func(
            f"=== Расчёт коэффициентов bk1, bk3, Bk (k={k}) при изменении {varied_param} ===")
        log_func(f"Диапазон: {start:.5f} → {end:.5f} | k = {k} | L = {L}\n")

    for val in values:
        bk1, bk3, Bk = zero_stability_coefficients_with_k(
            val, params, varied_param, k=k, L=L)
        if bk1 is not None:
            bk1_list.append(bk1)
            bk3_list.append(bk3)
            Bk_list.append(Bk)
            valid_values.append(val)

    if not valid_values:
        if log_func:
            log_func("Ошибка: не удалось найти точку P.")
        return

    # ====================== ПОСТРОЕНИЕ ГРАФИКА ======================
    if ax is None:
        fig, axs = plt.subplots(3, 1, figsize=(14, 10), sharex=True)
        external = True
    else:
        fig = ax.figure
        fig.clf()
        axs = fig.subplots(3, 1, sharex=True)
        external = False

    param_label = f'${varied_param}$'

    # --- Вертикальные линии начала и конца диапазона ---
    for axx in axs:
        axx.axvline(x=start, color='gray', linestyle='--',
                    linewidth=1.1, alpha=0.75)
        axx.axvline(x=end,   color='gray', linestyle='--',
                    linewidth=1.1, alpha=0.75)

    # --- Основные графики ---
    axs[0].plot(valid_values, bk1_list, 'b-', linewidth=2.3, label=f'k = {k}')
    axs[0].set_ylabel('$b_{k1}$', fontsize=14)
    axs[0].grid(True, alpha=0.35)
    axs[0].axhline(y=0, color='k', linestyle='--', linewidth=0.8)

    axs[1].plot(valid_values, bk3_list, 'g-', linewidth=2.3)
    axs[1].set_ylabel('$b_{k3}$', fontsize=14)
    axs[1].grid(True, alpha=0.35)
    axs[1].axhline(y=0, color='k', linestyle='--', linewidth=0.8)

    axs[2].plot(valid_values, Bk_list, 'r-', linewidth=2.3)
    axs[2].set_ylabel('$B_k = b_{k1} b_{k2} - b_{k3}$', fontsize=14)
    axs[2].set_xlabel(param_label, fontsize=14)
    axs[2].grid(True, alpha=0.35)
    axs[2].axhline(y=0, color='k', linestyle='--', linewidth=0.8)

    # ====================== РАЗМЕТКА start / end ======================
    # Подсвечиваем подписи на нижней оси
    current_ticks = axs[2].get_xticks()
    new_ticks = sorted(set(np.append(current_ticks, [start, end])))
    axs[2].set_xticks(new_ticks)

    labels = axs[2].get_xticklabels()
    for label in labels:
        try:
            val = float(label.get_text())
            if abs(val - start) < 1e-8 or abs(val - end) < 1e-8:
                label.set_fontweight('bold')
                label.set_color('darkred')
        except:
            pass

    # Красные маркеры на оси X (на нижнем графике)
    axs[2].plot([start, end], [0, 0], 'o', color='red', markersize=7,
                markeredgecolor='black', markeredgewidth=1.2, zorder=5, clip_on=False)

    # Заголовок
    fig.suptitle(
        f'Коэффициенты устойчивости для гармоники k = {k}\nв зависимости от параметра ${varied_param}$',
        fontsize=16)

    # чуть поднимаем, чтобы заголовок не налезал
    fig.tight_layout(rect=[0, 0, 1, 0.96])

    if external:
        plt.show()
    else:
        if canvas:
            canvas.draw()

    if log_func:
        log_func(f"Расчёт завершён для k={k} ({len(valid_values)} точек).")
